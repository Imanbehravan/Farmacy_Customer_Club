
from datetime import date, datetime

import jdatetime
from django.contrib.admin.views.decorators import staff_member_required
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.prescriptions.models import Prescription


# ---------------------------------------------------------------- کمکی‌ها
def to_jalali(value, with_time=False):
    """تبدیل date/datetime میلادی به رشته‌ی شمسی (به وقت تهران)."""
    if not value:
        return ""
    fmt = "%Y/%m/%d %H:%M" if with_time else "%Y/%m/%d"
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return jdatetime.datetime.fromgregorian(datetime=value).strftime(fmt)
    if isinstance(value, date):
        return jdatetime.date.fromgregorian(date=value).strftime("%Y/%m/%d")
    return str(value)


def parse_date(value):
    """
    تاریخ را از GET می‌خواند. هم شمسی (1405/01/15) و هم میلادی (2026-04-04) را قبول می‌کند.
    اگر سال کمتر از 1700 باشد شمسی فرض می‌شود. نامعتبر = None.
    """
    if not value:
        return None
    try:
        y, m, d = [int(x) for x in value.strip().replace("-", "/").split("/")]
        if y < 1700:
            return jdatetime.date(y, m, d).togregorian()
        return date(y, m, d)
    except (ValueError, TypeError):
        return None


def to_int(value):
    try:
        return int(value)
    except (ValueError, TypeError):
        return None


def display(obj, field):
    """اگر فیلد choices دارد برچسب خوانا (مثلاً «تحویل شده») را برمی‌گرداند، وگرنه خود مقدار."""
    getter = getattr(obj, f"get_{field}_display", None)
    value = getter() if getter else getattr(obj, field, "")
    return value if value is not None else ""


def get_filtered_prescriptions(request):
    """
    فیلترهای فرم گزارش: date_from, date_to, min_amount, max_amount, phone
    ⚠️ اگر صفحه‌ی گزارش شما فقط نسخه‌های خاصی را حساب می‌کند (مثلاً فقط status تحویل‌شده)،
    همان شرط را اینجا اضافه کنید، مثلاً: qs = qs.filter(status="delivered")
    """
    qs = Prescription.objects.select_related("user").order_by("-created_at")

    date_from = parse_date(request.GET.get("date_from"))
    date_to = parse_date(request.GET.get("date_to"))
    min_amount = to_int(request.GET.get("min_amount"))
    max_amount = to_int(request.GET.get("max_amount"))
    phone = (request.GET.get("phone") or "").strip()

    if date_from:
        qs = qs.filter(created_at__date__gte=date_from)
    if date_to:
        qs = qs.filter(created_at__date__lte=date_to)
    if min_amount is not None:
        qs = qs.filter(amount__gte=min_amount)
    if max_amount is not None:
        qs = qs.filter(amount__lte=max_amount)
    if phone:
        # در صفحه‌ی خطا، کاربر با شماره‌ی موبایل نمایش داده می‌شد، پس احتمالاً username همان شماره است
        qs = qs.filter(user__username__icontains=phone)
    return qs


# ---------------------------------------------------------------- ویو
@staff_member_required
def sales_report_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش فروش"
    ws.sheet_view.rightToLeft = True

    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F6F5C")

    headers = [
        "ردیف", "کد رهگیری", "مشتری", "کد ملی", "نوع بیمه",
        "مبلغ اصلی (ریال)", "مبلغ نهایی (ریال)", "وضعیت قرارداد", "وضعیت",
        "تاریخ ثبت", "تاریخ تحویل", "تاریخ عدم مراجعه", "یادداشت مدیر",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28

    row_num = 1
    for index, p in enumerate(get_filtered_prescriptions(request).iterator(), start=1):
        ws.append([
            index,
            p.tracking_code,
            str(p.user),
            p.national_id,
            display(p, "insurance_type"),
            p.original_amount,
            p.amount,
            display(p, "contract_status"),
            display(p, "status"),
            to_jalali(p.created_at),
            to_jalali(p.delivered_at),
            to_jalali(p.unclaimed_at),
            p.admin_notes or "",
        ])
        row_num = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=6).number_format = "#,##0"
        ws.cell(row=row_num, column=7).number_format = "#,##0"
        # کد ملی و کد رهگیری را متن نگه می‌داریم تا صفر ابتدایی حذف نشود
        ws.cell(row=row_num, column=2).number_format = "@"
        ws.cell(row=row_num, column=4).number_format = "@"

    # جمع کل با فرمول واقعی اکسل
    if row_num > 1:
        total_row = row_num + 1
        ws.cell(row=total_row, column=3, value="جمع کل")
        ws.cell(row=total_row, column=6, value=f"=SUM(F2:F{row_num})")
        ws.cell(row=total_row, column=7, value=f"=SUM(G2:G{row_num})")
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=total_row, column=col)
            cell.font = Font(bold=True)
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
        ws.cell(row=total_row, column=6).number_format = "#,##0"
        ws.cell(row=total_row, column=7).number_format = "#,##0"

    widths = [8, 16, 24, 16, 16, 18, 18, 16, 14, 14, 14, 16, 35]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    stamp = jdatetime.datetime.now().strftime("%Y%m%d")
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="sales_report_{stamp}.xlsx"'
    wb.save(response)
    return response